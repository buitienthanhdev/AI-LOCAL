"""Document Intelligence: trích văn bản THEO VỊ TRÍ (trang / mục / sheet+dòng / slide / dòng) + metadata.
Mỗi đoạn mang `loc` để câu trả lời trích dẫn được đến tận chỗ trong file gốc."""
import csv
import json
import logging
from pathlib import Path

logging.getLogger("pypdf").setLevel(logging.ERROR)  # PDF lỗi nhẹ rất phổ biến, không cần log

XLSX_BLOCK_ROWS = 25
LINE_BLOCK_CHARS = 1500


def _fmt(dt):
    return dt.isoformat(timespec="seconds") if dt else None


def _office_meta(cp) -> dict:
    # python-docx/pptx dùng author/last_modified_by; openpyxl dùng creator/lastModifiedBy
    return {
        "doc_author": getattr(cp, "author", None) or getattr(cp, "creator", None) or None,
        "doc_last_editor": getattr(cp, "last_modified_by", None) or getattr(cp, "lastModifiedBy", None) or None,
        "doc_created_at": _fmt(cp.created),
        "title": cp.title or None,
    }


def _line_blocks(text: str, prefix: str = "dòng") -> list[tuple[str, str]]:
    """Chia văn bản thường theo khối dòng, loc = 'dòng a–b'."""
    out, buf, start, n = [], [], 1, 0
    for i, line in enumerate(text.splitlines(), 1):
        buf.append(line)
        n += len(line) + 1
        if n >= LINE_BLOCK_CHARS:
            out.append((f"{prefix} {start}–{i}", "\n".join(buf)))
            buf, start, n = [], i + 1, 0
    if any(x.strip() for x in buf):
        out.append((f"{prefix} {start}–{start + len(buf) - 1}", "\n".join(buf)))
    return out


def _md_sections(text: str) -> list[tuple[str, str]]:
    out, path, buf = [], [], []

    def flush():
        body = "\n".join(buf).strip()
        if body:
            out.append(("mục: " + " › ".join(path) if path else "đầu tài liệu", body))
        buf.clear()
    for line in text.splitlines():
        if line.startswith("#"):
            flush()
            level = len(line) - len(line.lstrip("#"))
            path[:] = path[:level - 1] + [line.lstrip("# ").strip()]
        buf.append(line)
    flush()
    return out


def extract_segments(path: Path, meta_only: bool = False) -> tuple[list[tuple[str, str]], dict]:
    """(segments [(loc, text)], metadata). Lỗi sẽ được raise để nơi gọi ghi nhận.
    meta_only=True: chỉ lấy metadata nhanh, không dựng văn bản/không gọi model (dùng khi quét)."""
    ext = path.suffix.lower()
    from .media import MEDIA_EXT
    if ext in MEDIA_EXT:
        if meta_only:
            return [], {}
        from .media import read_media
        text = read_media(path)
        parts = [p for p in text.split("## ") if p.strip()]
        return [(("nội dung " + p.split("\n", 1)[0].strip().lower()), p.split("\n", 1)[-1])
                for p in parts] or [("nội dung", text)], {}

    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        meta = _office_meta(d.core_properties)
        if meta_only:
            return [], meta
        segs, heading, buf = [], [], []

        def flush():
            body = "\n".join(buf).strip()
            if body:
                segs.append(("mục: " + " › ".join(heading) if heading else "đầu tài liệu", body))
            buf.clear()
        for p in d.paragraphs:
            if not p.text.strip():
                continue
            style = (p.style.name if p.style is not None else None) or ""  # đoạn trỏ tới style không tồn tại => None
            if style.lower().startswith(("heading", "title")):
                flush()
                try:
                    lvl = max(1, int(style.split()[-1]))
                except (ValueError, IndexError):
                    lvl = 1
                heading[:] = heading[:lvl - 1] + [p.text.strip()]
            buf.append(p.text)
        flush()
        for i, t in enumerate(d.tables, 1):
            rows = [" | ".join(c.text.strip() for c in r.cells) for r in t.rows]
            if any(rows):
                segs.append((f"bảng {i}", "\n".join(rows)))
        return segs, meta

    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        meta = _office_meta(wb.properties)
        if meta_only:
            wb.close()
            return [], meta
        segs = []
        for ws in wb.worksheets:
            header, block, first = None, [], None
            for r, row in enumerate(ws.iter_rows(values_only=True), 1):
                cells = ["" if v is None else str(v) for v in row]
                if not any(cells):
                    continue
                line = " | ".join(cells)
                if header is None:
                    header = line
                if first is None:
                    first = r
                block.append(line)
                if len(block) >= XLSX_BLOCK_ROWS:
                    body = "\n".join(block if block[0] == header else [header, *block])
                    segs.append((f"sheet '{ws.title}' · dòng {first}–{r}", body))
                    block, first = [], None
            if block:
                body = "\n".join(block if block[0] == header else [header, *block])
                segs.append((f"sheet '{ws.title}' · dòng {first}–{first + len(block) - 1}", body))
        wb.close()
        return segs, meta

    if ext == ".pptx":
        import pptx
        prs = pptx.Presentation(str(path))
        meta = _office_meta(prs.core_properties)
        if meta_only:
            return [], meta
        segs = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                texts.append("Ghi chú: " + slide.notes_slide.notes_text_frame.text)
            if texts:
                segs.append((f"slide {i}", "\n".join(texts)))
        return segs, meta

    if ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(str(path))
        info = r.metadata or {}
        meta = {"doc_author": info.get("/Author"), "doc_last_editor": None,
                "doc_created_at": _fmt(getattr(info, "creation_date", None)), "title": info.get("/Title")}
        if meta_only:
            return [], meta
        return [(f"trang {i}", t) for i, p in enumerate(r.pages, 1) if (t := (p.extract_text() or "").strip())], meta

    if meta_only:
        return [], {}
    raw = path.read_text(encoding="utf-8", errors="replace")
    if ext == ".md":
        return _md_sections(raw), {}
    if ext == ".csv":
        rows = [" | ".join(r) for r in csv.reader(raw.splitlines())]
        header, segs = rows[0] if rows else "", []
        for i in range(0, len(rows), XLSX_BLOCK_ROWS):
            blk = rows[i:i + XLSX_BLOCK_ROWS]
            segs.append((f"dòng {i + 1}–{i + len(blk)}", "\n".join(blk if i == 0 else [header, *blk])))
        return segs, {}
    if ext == ".json":
        try:
            raw = json.dumps(json.loads(raw), ensure_ascii=False, indent=1)
        except ValueError:
            pass
    return _line_blocks(raw), {}


def extract(path: Path, light: bool = False) -> tuple[str, dict]:
    """Trả về (text, metadata) — văn bản phẳng (dùng cho trích tri thức, so sánh phiên bản).
    light=True: chỉ metadata."""
    segs, meta = extract_segments(path, meta_only=light)
    return "\n\n".join(t if loc.startswith("đầu") else f"[{loc}]\n{t}" for loc, t in segs), meta


def chunk(text: str, size: int, overlap: int) -> list[str]:
    text = text.strip()
    if not text:
        return []
    out, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):  # cắt ở xuống dòng gần nhất cho đẹp
            nl = text.rfind("\n", start + size // 2, end)
            if nl > 0:
                end = nl
        out.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return [c for c in out if c]


def chunk_segments(segments: list[tuple[str, str]], size: int, overlap: int, max_chunks: int) -> list[tuple[str, str]]:
    """Segments -> chunks [(loc, text)]. Đoạn dài bị cắt (loc thêm 'phần k'); đoạn quá ngắn liền kề được gộp."""
    out: list[tuple[str, str]] = []
    pend_loc, pend_txt = None, ""
    for loc, text in segments:
        text = text.strip()
        if not text:
            continue
        if len(text) < size // 4 and len(pend_txt) + len(text) < size:  # gộp đoạn ngắn
            pend_loc = (pend_loc + " + " + loc) if pend_loc else loc
            pend_txt = (pend_txt + "\n\n" if pend_txt else "") + text
            continue
        if pend_txt:
            out.append((pend_loc, pend_txt))
            pend_loc, pend_txt = None, ""
        parts = chunk(text, size, overlap)
        for k, part in enumerate(parts, 1):
            out.append((loc if len(parts) == 1 else f"{loc} (phần {k}/{len(parts)})", part))
    if pend_txt:
        out.append((pend_loc, pend_txt))
    return out[:max_chunks]
