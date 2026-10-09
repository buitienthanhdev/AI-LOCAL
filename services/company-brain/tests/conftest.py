"""Hạ tầng test: dữ liệu RIÊNG trong thư mục tạm (không đụng dữ liệu thật), Qdrant trong RAM, embedding giả xác định
(chỉ trong test — production luôn dùng model thật). Test đánh dấu @pytest.mark.live cần Ollama thật: chạy với --live."""
import hashlib
import math
import os
import re
import shutil
import sys
import tempfile
from pathlib import Path

import pytest

_ROOT = Path(__file__).resolve().parent.parent
# Không dùng %TEMP% (nằm trong AppData => hệ thống coi là thư mục phần mềm: bậc 3, watcher bỏ qua)
(_ROOT / ".test_tmp").mkdir(exist_ok=True)
_TMP = Path(tempfile.mkdtemp(prefix="run_", dir=_ROOT / ".test_tmp"))
os.environ["BRAIN_DATA_DIR"] = str(_TMP / "data")
os.environ["BRAIN_QDRANT"] = "memory"
sys.path.insert(0, str(_ROOT))


def pytest_addoption(parser):
    parser.addoption("--live", action="store_true", help="chạy cả test cần Ollama/model thật")


def pytest_collection_modifyitems(config, items):
    if not config.getoption("--live"):
        skip = pytest.mark.skip(reason="cần --live (Ollama + model thật)")
        for it in items:
            if "live" in it.keywords:
                it.add_marker(skip)


def fake_embed(texts, query=False):
    """Embedding túi-từ băm 128 chiều (xác định, không cần GPU) — CHỈ dùng trong test."""
    from brain.facts import norm
    out = []
    for t in texts:
        v = [0.0] * 128
        for w in re.findall(r"[a-z0-9.]+", norm(t)):
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % 128] += 1.0
        n = math.sqrt(sum(x * x for x in v)) or 1.0
        out.append([x / n for x in v])
    return out


@pytest.fixture(scope="session", autouse=True)
def _cleanup_tmp():
    yield
    shutil.rmtree(_TMP, ignore_errors=True)


def _safe(fn):
    try:
        fn()
    except Exception:
        pass


@pytest.fixture
def env(request, monkeypatch):
    """Môi trường sạch cho mỗi test: DB mới, kho vector mới, thư mục chia sẻ riêng."""
    import sqlite3
    from brain import catalog, config, index, scanner, watcher
    # theo dõi mọi kết nối SQLite mở trong test để đóng sạch khi test kết thúc (test hỏng không được để lại khóa ghi)
    opened: list = []
    real_connect = sqlite3.connect

    def tracking_connect(*a, **k):
        conn = real_connect(*a, **k)
        opened.append(conn)
        return conn
    monkeypatch.setattr(sqlite3, "connect", tracking_connect)
    request.addfinalizer(lambda: [(_safe(c.rollback), _safe(c.close)) for c in opened])
    data = Path(os.environ["BRAIN_DATA_DIR"])
    data.mkdir(parents=True, exist_ok=True)
    c = catalog.connect()  # xóa dữ liệu theo bảng (xóa file DB sẽ vướng khóa của luồng nền trên Windows)
    for (t,) in c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' "
                          "AND name NOT LIKE 'chunks_fts_%'").fetchall():
        if t != "chunks_fts":
            c.execute(f"DELETE FROM {t}")
    c.execute("INSERT INTO chunks_fts(chunks_fts) VALUES ('rebuild')")
    c.commit()
    c.close()
    shutil.rmtree(data / "versions", ignore_errors=True)
    shutil.rmtree(data / "datasets", ignore_errors=True)
    index._client = None
    # mỗi test một thư mục riêng (xóa-tạo lại hay vướng khóa tạm của Windows; dọn cả cây khi kết thúc phiên test)
    shared = _TMP / f"shared_{re.sub(r'[^a-z0-9]', '_', request.node.name.lower())[:60]}"
    shared.mkdir(parents=True, exist_ok=True)
    config.SHARED_ROOTS[:] = [shared]
    config.FACT_ROOTS[:] = [shared]
    config.ACL_RULES.clear()
    config.ACL_RULES.update({"HR": ["hr", "board"]})
    monkeypatch.setattr(watcher, "DEBOUNCE_SEC", 0)
    if "live" not in request.keywords:
        monkeypatch.setattr(index, "embed", fake_embed)
    from brain import registry
    registry._cache["at"] = 0
    yield shared
    index._client = None


def write_docx(path: Path, author: str, paragraphs: list[str], editor: str | None = None, headings: bool = False):
    import docx
    path.parent.mkdir(parents=True, exist_ok=True)
    d = docx.Document()
    for i, p in enumerate(paragraphs):
        if headings and i == 0:
            d.add_heading(p, 1)
        else:
            d.add_paragraph(p)
    d.core_properties.author = author
    d.core_properties.last_modified_by = editor or author
    d.save(str(path))
    return path


def scan_shared(shared: Path):
    from brain import scanner
    from brain.catalog import connect
    stats = scanner.scan(shared)
    c = connect()
    scanner.backfill_tiers(c, progress=lambda *_: None)
    c.close()
    return stats


def file_row(path: Path):
    from brain.catalog import connect
    c = connect()
    r = c.execute("SELECT * FROM files WHERE path=?", (str(path.resolve()),)).fetchone()
    c.close()
    return r
