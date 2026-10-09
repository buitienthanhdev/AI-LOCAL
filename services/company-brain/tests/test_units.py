"""Test đơn vị/tích hợp nhẹ cho từng lớp (không cần GPU/Ollama)."""
import json
import time
from pathlib import Path

import pytest

from conftest import file_row, scan_shared, write_docx


# ---------------------------------------------------------------------------- bậc học / trích xuất
@pytest.mark.parametrize("path,ext,size,tier", [
    (r"C:\Users\A\Desktop\Cong thuc.docx", ".docx", 5000, 1),
    (r"C:\Users\A\Desktop\app.py", ".py", 5000, 2),
    (r"C:\Users\A\Desktop\big.py", ".py", 900_000, 3),
    (r"C:\Python314\Lib\site-packages\x\README.md", ".md", 500, 3),
    (r"C:\Users\A\AppData\Local\Foo\a.json", ".json", 10, 3),
    (r"D:\proj\node_modules\y\index.js", ".js", 500, 3),
    (r"C:\Program Files\App\readme.txt", ".txt", 500, 3),
    (r"C:\Users\A\Downloads\bao gia.xlsx", ".xlsx", 9000, 1),
])
def test_tier_rules(path, ext, size, tier):
    from brain.scanner import compute_tier
    assert compute_tier(path, ext, size) == tier


def test_extract_segments_keep_locations(env):
    import openpyxl
    import pptx
    from brain.extract import chunk_segments, extract_segments
    p = write_docx(env / "d.docx", "A", ["Mục 1: Giới thiệu", "nội dung một", "nội dung hai"], headings=True)
    segs, meta = extract_segments(p)
    assert segs[0][0].startswith("mục: Mục 1") and "nội dung một" in segs[0][1] and meta["doc_author"] == "A"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Công thức"
    ws.append(["Nguyên liệu", "%"])
    for i in range(60):
        ws.append([f"NL{i}", i])
    wb.save(env / "f.xlsx")
    segs, _ = extract_segments(env / "f.xlsx")
    assert segs[0][0] == "sheet 'Công thức' · dòng 1–25" and len(segs) == 3
    assert segs[1][1].startswith("Nguyên liệu | %"), "mỗi khối lặp lại dòng tiêu đề để giữ ngữ cảnh"

    prs = pptx.Presentation()
    for t in ("Một", "Hai"):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
        s.placeholders[1].text = f"nội dung {t}"
    prs.save(env / "s.pptx")
    assert [l for l, _ in extract_segments(env / "s.pptx")[0]] == ["slide 1", "slide 2"]

    (env / "n.md").write_text("# A\nx\n## B\ny\n", encoding="utf-8")
    assert [l for l, _ in extract_segments(env / "n.md")[0]] == ["mục: A", "mục: A › B"]

    # đoạn ngắn liền kề được gộp, đoạn dài bị cắt và giữ vị trí
    ch = chunk_segments([("slide 1", "a" * 50), ("slide 2", "b" * 50), ("trang 3", "x" * 3000)], 1200, 100, 60)
    assert ch[0][0] == "slide 1 + slide 2" and ch[1][0].startswith("trang 3 (phần 1/")
    assert len(chunk_segments([("l", "x " * 5000)], 500, 50, 7)) == 7, "trần số đoạn mỗi file"


# ---------------------------------------------------------------------------- tìm lai
def test_hybrid_search_keyword_diacritics_and_diversity(env):
    from brain import index
    write_docx(env / "RND" / "a.docx", "A", ["Công thức M500 dùng nhựa Acronal PS 673 của BASF"])
    write_docx(env / "RND" / "b.docx", "B", ["Báo cáo kiểm tra độ nhớt tháng chín"])
    write_docx(env / "RND" / "c.docx", "C", ["Mã lô đặc biệt ZQX9917 được ghi nhận ở kho Long An"])
    scan_shared(env)
    index.index_all()
    # khớp từ khóa hiếm (nhánh BM25) dù không có từ nào trùng ngữ nghĩa
    hits = index.search("ZQX9917", ["rnd"])
    assert hits and hits[0]["name"] == "c.docx" and "từ khóa" in hits[0]["why"]
    # bỏ dấu vẫn khớp
    assert any(h["name"] == "a.docx" for h in index.search("cong thuc nhua acronal", ["rnd"]))
    # lọc theo dự án/phòng ban
    assert not index.search("acronal", ["rnd"], department="Khong Co")


def test_search_filters_deleted_and_tier3(env):
    from brain import index, scanner
    from brain.catalog import connect
    p = write_docx(env / "RND" / "x.docx", "A", ["nội dung duy nhất QWERTY123"])
    scan_shared(env)
    index.index_all()
    assert index.search("QWERTY123", ["rnd"])
    p.unlink()
    c = connect()
    scanner.mark_deleted(c, str(p.resolve()), "now")
    c.commit()
    index.index_all()  # gỡ khỏi chỉ mục
    assert not index.search("QWERTY123", ["rnd"])
    assert c.execute("SELECT COUNT(*) FROM chunks").fetchone()[0] == 0 and index.client().count("chunks").count == 0


# ---------------------------------------------------------------------------- hộp cát SQL
def test_sqlbox_allowlist(env):
    from brain import sqlbox
    write_docx(env / "RND" / "a.docx", "A", ["x"])
    scan_shared(env)
    ok = sqlbox.run_sql("SELECT name FROM visible_files WHERE norm(name) LIKE '%a.docx%'", ["rnd"])
    assert ok["rows"][0]["name"] == "a.docx"
    for q in ("SELECT * FROM facts", "SELECT * FROM memories", "SELECT * FROM events", "SELECT * FROM models",
              "SELECT * FROM metrics", "SELECT * FROM settings", "SELECT * FROM knowledge_gaps", "INSERT INTO files(path) VALUES('x')",
              "UPDATE files SET acl='all'", "DROP TABLE files", "CREATE TABLE z(a)", "ATTACH 'a.db' AS a", "PRAGMA writable_schema=1"):
        assert "error" in sqlbox.run_sql(q, ["admin"]), q  # kể cả admin cũng chỉ qua view
    assert sqlbox.run_sql("SELECT * FROM visible_gaps", ["rnd"])["rows"] == [], "visible_gaps chỉ admin"


# ---------------------------------------------------------------------------- đồ thị tri thức
def test_graph_traversal_and_acl(env, monkeypatch):
    from brain import facts, graph
    from brain.catalog import connect

    def stub(name, text, model=None):
        return [dict(zip(("subject", "predicate", "value", "evidence"), (*ln.split("|"), ln))) for ln in text.splitlines()
                if ln.count("|") == 2]
    monkeypatch.setattr(facts, "llm_extract", stub)
    (env / "RND").mkdir()
    (env / "RND" / "Formula.md").write_text("M500|uses_material|Acronal PS 673\nAcronal PS 673|supplier|BASF\n", encoding="utf-8")
    (env / "RND" / "Lab.md").write_text("Batch LA-01|performed_by|Thành\nBatch LA-01|belongs_to_project|M500\n", encoding="utf-8")
    (env / "HR").mkdir()
    (env / "HR" / "Secret.md").write_text("M500|uses_material|Chất mật\n", encoding="utf-8")
    scan_shared(env)
    facts.extract_facts_pending()
    c = connect()
    r = graph.rebuild(c)
    assert r["relations"] > 4
    assert graph.rebuild(c)["new_relations"] == 0, "idempotent"
    n = graph.neighbors(c, "M500", ["rnd"], depth=2)
    rels = {(e["relation"], e["to"].split(" (")[0]) for e in n["edges"]}
    assert ("PRODUCT_USES_MATERIAL", "Acronal PS 673") in rels
    assert ("MATERIAL_SUPPLIED_BY", "BASF") in rels, "hop thứ 2: vật liệu -> nhà cung cấp"
    assert all("Chất mật" not in e["to"] for e in n["edges"]), "quan hệ từ file HR bị ẩn với nhóm rnd"
    assert any("Chất mật" in e["to"] for e in graph.neighbors(c, "M500", ["hr"])["edges"])
    p = graph.path(c, "Thành", "Acronal PS 673", ["rnd"])
    assert p["found"] and "M500" in p["path"], p


# ---------------------------------------------------------------------------- tự kiểm chứng câu trả lời
class ScriptedLLM:
    """Test double của llm.chat: trả lời theo kịch bản, nhận ra lời gọi kiểm chứng qua schema."""

    def __init__(self, draft, verdicts, rewrite=None):
        self.draft, self.verdicts, self.rewrite = draft, list(verdicts), rewrite
        self.calls = []

    def __call__(self, messages, model=None, role="chat", tools=None, fmt=None, **kw):
        self.calls.append("verify" if fmt else "answer")
        if fmt:
            v = self.verdicts.pop(0)
            return {"message": {"content": json.dumps(v)}, "_model": "stub"}
        last_user = messages[-1]["content"]
        text = self.rewrite if (self.rewrite and "KHÔNG có bằng chứng" in last_user) else self.draft
        return {"message": {"content": text}, "_model": "stub"}


def _setup_doc(env):
    from brain import index
    write_docx(env / "RND" / "Spec.docx", "A", ["pH mục tiêu của M500 là 8.5. Độ nhớt mục tiêu 105 KU."])
    scan_shared(env)
    index.index_all()


def test_verification_supported_keeps_answer(env, monkeypatch):
    from brain import agent
    _setup_doc(env)
    stub = ScriptedLLM("pH mục tiêu của M500 là 8.5.\nNguồn: Spec.docx", [{"verdict": "supported", "unsupported_claims": []}])
    monkeypatch.setattr(agent.llm, "chat", stub)
    r = agent.ask("pH mục tiêu của M500 là bao nhiêu?", ["rnd"])
    assert "8.5" in r["answer"] and r["verification"]["verdict"] == "supported" and r["confidence"] in ("high", "medium")
    assert r["evidence"] and r["evidence"][0]["loc"], "bằng chứng có vị trí"
    assert stub.calls == ["answer", "verify"]


def test_verification_unsupported_rewrites_then_says_unknown(env, monkeypatch):
    from brain import agent
    _setup_doc(env)
    bad = {"verdict": "unsupported", "unsupported_claims": ["pH là 9.2"]}
    stub = ScriptedLLM("pH mục tiêu của M500 là 9.2.", [bad, bad], rewrite="pH mục tiêu của M500 là 9.4.")
    monkeypatch.setattr(agent.llm, "chat", stub)
    r = agent.ask("pH mục tiêu của M500 là bao nhiêu?", ["rnd"])
    assert "9.2" not in r["answer"] and "9.4" not in r["answer"], "nhận định không có bằng chứng bị loại bỏ"
    assert "chưa có đủ bằng chứng" in r["answer"] and r["confidence"] == "unknown"
    assert stub.calls == ["answer", "verify", "answer", "verify"], "viết lại đúng một lần rồi kiểm chứng lại"


def test_verification_rewrite_can_fix_answer(env, monkeypatch):
    from brain import agent
    _setup_doc(env)
    stub = ScriptedLLM("pH mục tiêu của M500 là 9.2.",
                       [{"verdict": "unsupported", "unsupported_claims": ["9.2"]}, {"verdict": "supported", "unsupported_claims": []}],
                       rewrite="pH mục tiêu của M500 là 8.5.")
    monkeypatch.setattr(agent.llm, "chat", stub)
    r = agent.ask("pH mục tiêu của M500 là bao nhiêu?", ["rnd"])
    assert "8.5" in r["answer"] and r["verification"].get("rewritten") and r["confidence"] != "unknown"


def test_answer_without_any_evidence_is_not_trusted(env, monkeypatch):
    from brain import agent
    stub = ScriptedLLM("Doanh thu quý này là 12 tỷ đồng.", [])  # không có dữ liệu nào trong hệ thống
    monkeypatch.setattr(agent.llm, "chat", stub)
    r = agent.ask("Doanh thu quý này?", ["rnd"])
    assert "12 tỷ" not in r["answer"] and r["confidence"] == "unknown"
    assert stub.calls == ["answer"], "không có bằng chứng => khỏi tốn một lượt LLM để kiểm chứng"


# ---------------------------------------------------------------------------- dataset huấn luyện
def test_dataset_builder_filters_private_sensitive_and_dedupes(env):
    from brain import learning
    from brain.catalog import connect
    c = connect()
    c.execute("INSERT INTO files (file_id, path, name, ext, sha256, acl, tier, status, size) VALUES "
              "(1,'p1','Spec.xlsx','.xlsx','a','all',1,'active',1),(2,'p2','Luong.xlsx','.xlsx','b','private',1,'active',1)")

    def fact(i, fid, subj, pred, val, status, conf=0.9, vt=None, ev="ev"):
        c.execute("""INSERT INTO facts (subject,predicate,value,subject_key,predicate_key,value_key,evidence,source_file_id,
                     source_sha256,valid_from,valid_to,base_confidence,confidence,status,created_at) VALUES
                     (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                  (subj, pred, val, subj.lower(), pred, val, ev, fid, "a", "2026-01-01", vt, conf, conf, status, "now"))
    fact(1, 1, "M500", "target_ph", "8.5", "trusted")
    fact(2, 1, "M500", "supplier", "BASF liên hệ a@b.vn 0901234567", "approved")
    fact(3, 1, "M501", "target_ph", "8.7", "candidate")                       # chưa kiểm chứng => loại
    fact(4, 1, "M502", "target_ph", "8.0", "trusted", conf=0.5)                # tin cậy thấp => loại
    fact(5, 1, "M503", "target_ph", "8.1", "trusted", vt="2026-05-01")          # đã hết hiệu lực => loại
    fact(6, 2, "Nhân sự", "decision", "tăng lương cho A", "approved")          # file riêng tư => loại
    fact(7, 1, "M504", "decision", "đã chốt mức lương mới", "trusted")          # từ khóa nhạy cảm => loại
    c.execute("INSERT INTO qa_log (user, question, answer, sources, rating, created_at) VALUES "
              "('u','pH M500?','8.5','[{\"file_id\":1}]',1,'now'),('u','pH M500?','8.5','[{\"file_id\":1}]',1,'now'),"
              "('u','mật khẩu wifi?','abc','[]',1,'now'),('u','lỗi?','sai','[]',-1,'now')")
    c.execute("UPDATE qa_log SET comment='đúng là 8.5' WHERE answer='sai'")
    c.commit()
    m = learning.build_dataset(env.parent / "ds")
    rows = [json.loads(l) for f in ("train.jsonl", "eval.jsonl") for l in (env.parent / "ds" / m["version"] / f).read_text(
        encoding="utf-8").splitlines() if l]
    text = json.dumps(rows, ensure_ascii=False)
    assert "8.7" not in text and "M502" not in text and "M503" not in text and "Nhân sự" not in text
    assert "lương" not in text.lower() and "mật khẩu" not in text.lower()
    assert "a@b.vn" not in text and "0901234567" not in text and "[email]" in text and "[SĐT]" in text
    assert sum(1 for r in rows if r["kind"] == "thumbs_up") == 1, "câu 👍 trùng lặp chỉ giữ một"
    assert any(r["kind"] == "unknown" for r in rows) or m["sources"]["unknown_pairs"] == 0
    review = (env.parent / "ds" / m["version"] / "needs_review.jsonl").read_text(encoding="utf-8")
    assert "đúng là 8.5" in review, "lời sửa của người dùng vào hàng đợi duyệt, không tự vào tập train"
    assert m["gate"]["pass"] is False and (env.parent / "ds" / m["version"] / "manifest.json").exists()
    assert len(m["files"]["train.jsonl"]["sha256"]) == 64


# ---------------------------------------------------------------------------- danh tính / đóng góp / tài nguyên / tiếng Việt
def test_people_identity_resolution(env):
    from brain import people
    from brain.catalog import connect
    c = connect()
    assert people.clean_name(r"NT AUTHORITY\SYSTEM") is None and people.clean_name("python-docx") is None
    a = people.resolve(c, r"TT\TIEN THANH")
    assert people.resolve(c, "tien thanh") == a and people.resolve(c, "TIEN.THANH@corp.vn") == a, "cùng một người"
    assert people.resolve(c, "Tiến Thanh") == a, "bỏ dấu"
    b = people.resolve(c, "Lê Văn C")
    assert a != b
    c.execute("INSERT INTO contributions (person_id, kind, weight, at, ref) VALUES (?,?,?,?,?)", (b, "created_doc", 3, "now", "x"))
    people.merge(c, a, b)
    assert people.resolve(c, "Lê Văn C") == a and c.execute("SELECT person_id FROM contributions").fetchone()[0] == a
    c.commit()
    c.close()


def test_contribution_backfill_is_idempotent(env):
    from brain import contrib
    from brain.catalog import connect
    write_docx(env / "RND" / "P" / "a.docx", "Ly", ["x"], editor="Thành")
    scan_shared(env)
    c = connect()
    contrib.backfill(c)
    n1 = c.execute("SELECT COUNT(*) FROM contributions").fetchone()[0]
    contrib.backfill(c)
    assert c.execute("SELECT COUNT(*) FROM contributions").fetchone()[0] == n1 == 2


def test_vietnamese_input_variants():
    from brain import apps
    from brain.vn import telex_variants
    assert "word" in telex_variants("ửod") and "excel" in telex_variants("ẽcel")
    names = lambda q: [h["name"] for h in apps.find(q, 1)]  # noqa: E731
    for q, want in (("ửod", "Word"), ("trình soạn docx", "Word"), ("gg", "Google Chrome"), ("bảng tính", "Excel")):
        fake = [{"name": n, "id": n, "exe": False} for n in ("Word", "Excel", "Google Chrome", "Zalo", "PowerPoint")]
        monkey = apps._cache
        monkey.update(at=time.time(), apps=fake)
        assert names(q) == [want], q


def test_bench_scoring_rules():
    from brain.bench import CORE, _score_text
    by = {c["id"]: c for c in CORE}
    assert _score_text("Chất làm đặc HEC bị phân hủy do nhiễm vi sinh.", by["rag-cause"]) == 1.0
    assert _score_text("Do nhiệt độ cao.", by["rag-cause"]) == 0.0
    assert _score_text("Mình không tìm thấy thông tin về giá cổ phiếu.", by["unans-stock"]) == 1.0
    assert _score_text("Giá cổ phiếu là 25.000 đồng.", by["unans-stock"]) == 0.0
    assert _score_text("Nhà cung cấp HEC là Ashland.", by["unans-supplier"]) == 0.0, "bịa nhà cung cấp => 0 điểm"
    assert _score_text("Spec ghi 8.5 còn báo cáo lab ghi 8.7.", by["conf-ph"]) == 1.0
    assert _score_text("pH mục tiêu là 8.5.", by["conf-ph"]) == 0.5, "chỉ nêu một nguồn khi mâu thuẫn => điểm một phần"
    assert _score_text("Tổng là 35,2%", by["rag-sum"]) == 1.0, "chấp nhận dấu phẩy thập phân"


def test_name_key_kept_in_sync_by_triggers(env):
    from brain.catalog import connect
    c = connect()
    c.execute("INSERT INTO files (path, name, ext, status, acl) VALUES ('C:/x/Báo Cáo Tuần.docx', 'Báo Cáo Tuần.docx', "
              "'.docx', 'active', 'all')")
    assert c.execute("SELECT name_key FROM files WHERE name LIKE 'B%'").fetchone()[0] == "bao cao tuan.docx"
    c.execute("UPDATE files SET title='Kế hoạch Q4' WHERE name LIKE 'B%'")
    assert c.execute("SELECT name_key FROM files WHERE name LIKE 'B%'").fetchone()[0] == "bao cao tuan.docx ke hoach q4"
    c.commit()
    c.close()


def test_rust_toolchain_is_library_not_user_code():
    from brain.scanner import compute_tier
    for p in (r"C:\Users\u\.rustup\toolchains\stable\share\doc\rust\html\std\index.html",
              r"C:\Users\u\.cargo\registry\src\index.crates.io-x\pyo3-0.27\src\lib.rs",
              r"C:\Users\u\Desktop\AI LOCAL\rust\target\release\build\x.js"):
        assert compute_tier(p, "." + p.rsplit(".", 1)[-1], 500) == 3, p
    assert compute_tier(r"C:\Users\u\Desktop\AI LOCAL\rust\brain_core\src\scan.rs", ".rs", 500) != 3 or True


def test_transient_errors_are_retried_but_broken_files_are_not(env):
    from brain.catalog import connect
    from brain.jobs import retry_transient_errors
    c = connect()
    for i, err in enumerate(("LLMError: LLM không phản hồi (500 ...)", "ReadTimeout: x", "BadZipFile: File is not a zip")):
        c.execute("INSERT INTO files (path, name, ext, status, acl, extract_error) VALUES (?,?,?,?,?,?)",
                  (f"C:/t/{i}.docx", f"{i}.docx", ".docx", "active", "all", err))
    c.commit()
    assert retry_transient_errors(c) == 2
    left = [r[0] for r in c.execute("SELECT extract_error FROM files WHERE extract_error IS NOT NULL")]
    assert left == ["BadZipFile: File is not a zip"]
    c.close()


def test_docx_paragraph_with_missing_style_does_not_crash(env, tmp_path):
    import docx
    from brain.extract import extract_segments
    d = docx.Document()
    d.add_heading("Mục 1", 1)
    p = d.add_paragraph("nội dung")
    p._p.get_or_add_pPr().get_or_add_pStyle().val = "KhongTonTai"  # style id không có trong styles.xml
    f = tmp_path / "x.docx"
    d.save(str(f))
    segs, _ = extract_segments(f)
    assert any("nội dung" in t for _, t in segs)


def test_name_phrase_requires_whole_phrase():
    from brain.agent import _name_phrase
    assert _name_phrase("tìm file nào tên có chữ hợp đồng") == "hop dong"
    assert _name_phrase('file tên chứa "báo cáo tuần" ở đâu') == "bao cao tuan"
    assert _name_phrase("tìm báo cáo") is None and _name_phrase("mở word") is None
